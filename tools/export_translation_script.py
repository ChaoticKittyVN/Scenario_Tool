"""
导出翻译表格工具
从 Excel 剧本中提取角色名和文本，生成供翻译填写的表格
"""
import argparse
import sys
from pathlib import Path
from typing import Optional

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from core.config_manager import AppConfig
from core.constants import ColumnName, SheetName, SpecialName
from core.excel_management.dataframe_processor import DataFrameProcessor
from core.excel_management.excel_file_manager import (
    ExcelFileManager,
    ExcelFileNotFoundError,
    ExcelFormatError,
)
from core.logger import get_logger
from core.param_process.param_translator import ParamTranslator

logger = get_logger(__name__)

SPECIAL_NAME_VALUES = {member.value for member in SpecialName}
OUTPUT_COLUMNS = [
    "Excel",
    "Sheet",
    "Idx",
    "Index",
    "Name",
    "Text",
    "TranslatedName",
    "TranslatedText",
    "TextID",
]


class TranslationTableExporter:
    """翻译表格导出器。"""

    COLUMN_WIDTHS = {
        "Excel": 20,
        "Sheet": 14,
        "Idx": 10,
        "Index": 12,
        "Name": 14,
        "Text": 60,
        "TranslatedName": 18,
        "TranslatedText": 60,
        "TextID": 10,
    }

    def __init__(
        self,
        config: AppConfig,
        input_dir: Optional[Path] = None,
        output_dir: Optional[Path] = None,
        merge_files: bool = False,
        output_format: str = "excel",
    ):
        self.input_dir = input_dir or config.paths.input_dir
        self.output_dir = output_dir or (config.paths.output_dir / "translation_exports")
        self.merge_files = merge_files
        self.output_format = output_format.lower()
        self.ignore_words = {"无语音"}

        self.excel_manager = ExcelFileManager(cache_enabled=True)
        self.df_processor = DataFrameProcessor(config)
        self.translator = ParamTranslator()

    def extract_from_file(self, file_path: Path) -> pd.DataFrame:
        """从单个 Excel 文件提取待翻译内容。"""
        logger.info(f"正在处理：{file_path.name}")
        try:
            excel_data = self.excel_manager.load_excel(file_path)
        except (ExcelFileNotFoundError, ExcelFormatError) as exc:
            logger.error(f"跳过文件 {file_path.name}: {exc}")
            return pd.DataFrame(columns=OUTPUT_COLUMNS)

        rows = []
        for sheet_name, df in excel_data.items():
            if sheet_name == SheetName.PARAM_SHEET.value:
                continue

            valid_df = self.df_processor.extract_valid_rows(df, sheet_name)
            if valid_df.empty:
                continue

            for idx, row in valid_df.iterrows():
                ignore = row.get(ColumnName.IGNORE.value, "")
                if not pd.isna(ignore) and ignore in self.ignore_words:
                    continue

                text = row.get(ColumnName.TEXT.value, "")
                if pd.isna(text) or str(text).strip() == "":
                    continue

                speaker = row.get(ColumnName.NAME.value, "")
                speaker = "" if pd.isna(speaker) else str(speaker).strip()
                if speaker in SPECIAL_NAME_VALUES and not speaker in [SpecialName.TEXT_COMMAND.value, SpecialName.NVL_COMMAND.value]:
                    continue

                line_index = row.get(ColumnName.INDEX.value, "")
                line_index = "" if pd.isna(line_index) else line_index

                translated_speaker = ""
                if speaker and self.translator.has_mapping("NameTranslated", speaker):
                    translated_speaker = self.translator.translate("NameTranslated", speaker)
                else:
                    translated_speaker = speaker

                rows.append({
                    "Excel": file_path.stem,
                    "Sheet": sheet_name,
                    "Idx": idx + 2,
                    "Index": line_index,
                    "Name": speaker,
                    "Text": str(text).strip(),
                    "TranslatedName": translated_speaker,
                    "TranslatedText": "",
                    "TextID": "",
                })

        result = pd.DataFrame(rows, columns=OUTPUT_COLUMNS)
        logger.info(f"  提取到 {len(result)} 条文本")
        return result

    def export(self):
        """执行导出。"""
        excel_files = sorted(
            list(self.input_dir.glob("*.xlsx")) + list(self.input_dir.glob("*.xls"))
        )
        excel_files = [file for file in excel_files if not file.name.startswith("~")]

        if not excel_files:
            logger.warning(f"在 {self.input_dir} 中没有找到 Excel 文件")
            return

        self.output_dir.mkdir(parents=True, exist_ok=True)

        if self.merge_files:
            dataframes = [self.extract_from_file(file) for file in excel_files]
            dataframes = [df for df in dataframes if not df.empty]
            if not dataframes:
                logger.warning("没有提取到任何文本")
                return
            combined_df = pd.concat(dataframes, ignore_index=True)
            self._save_dataframe(combined_df, self.output_dir / "全部翻译表格")
            return

        for file_path in excel_files:
            dataframe = self.extract_from_file(file_path)
            if dataframe.empty:
                continue
            self._save_dataframe(dataframe, self.output_dir / f"{file_path.stem}_翻译表格")

    def _save_dataframe(self, df: pd.DataFrame, base_path: Path):
        """按原工作表结构保存翻译表格。"""
        if self.output_format == "excel":
            output_path = base_path.with_suffix(".xlsx")
            with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
                for sheet_name, sheet_df in df.groupby("Sheet", sort=False):
                    sheet_df.to_excel(writer, index=False, sheet_name=str(sheet_name))
                    worksheet = writer.sheets[str(sheet_name)]
                    center_alignment = Alignment(horizontal="center", vertical="center")
                    text_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

                    for column_index, column_name in enumerate(sheet_df.columns, start=1):
                        column_letter = get_column_letter(column_index)
                        worksheet.column_dimensions[column_letter].width = self.COLUMN_WIDTHS[column_name]

                        worksheet.cell(row=1, column=column_index).alignment = center_alignment
                        data_alignment = (
                            text_alignment
                            if column_name in {"Text", "TranslatedText", "TextID"}
                            else center_alignment
                        )
                        for row_index in range(2, worksheet.max_row + 1):
                            worksheet.cell(row=row_index, column=column_index).alignment = data_alignment
        else:
            output_path = base_path.with_suffix(".csv")
            df.to_csv(output_path, index=False, encoding="utf-8-sig")

        logger.info(f"已导出 {len(df)} 行到 {output_path}")


def main():
    parser = argparse.ArgumentParser(description="导出翻译表格工具")
    parser.add_argument("--input", "-i", type=Path, help="输入目录（覆盖配置中的 input_dir）")
    parser.add_argument("--output", "-o", type=Path, help="输出目录（覆盖配置中的 output_dir）")
    parser.add_argument("--merge", action="store_true", help="合并所有文件为一个翻译表格")
    parser.add_argument(
        "--format",
        choices=["excel", "csv"],
        default="excel",
        help="输出格式（默认：excel）",
    )
    args = parser.parse_args()

    config_path = Path("config.yaml")
    if config_path.exists():
        config = AppConfig.from_file(config_path)
    else:
        logger.warning("config.yaml 不存在，使用默认配置")
        config = AppConfig.create_default("naninovel")

    exporter = TranslationTableExporter(
        config=config,
        input_dir=args.input,
        output_dir=args.output,
        merge_files=args.merge,
        output_format=args.format,
    )
    exporter.export()


if __name__ == "__main__":
    main()
