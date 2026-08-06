from pathlib import Path

import pandas as pd
import pytest
import yaml
from openpyxl import load_workbook

from core.table_transform import TableTransformEngine, load_definition
from core.table_transform.conditions import matches_condition


def write_config(path: Path, operations, conflict_policy="error") -> Path:
    path.write_text(
        yaml.safe_dump(
            {
                "version": 1,
                "conflict_policy": conflict_policy,
                "operations": operations,
            },
            allow_unicode=True,
            sort_keys=False,
        ),
        encoding="utf-8",
    )
    return path


def write_workbook(path: Path) -> Path:
    pd.DataFrame(
        [
            {
                "Name": "Alice",
                "Text": "Hello",
                "Background": "room",
                "Position": "left",
                "Voice": "",
                "Volume": "",
                "Index": "",
                "Source": "A",
                "Mapped": "",
            },
            {
                "Name": "Alice",
                "Text": "Again",
                "Background": "",
                "Position": "",
                "Voice": "",
                "Volume": "",
                "Index": "",
                "Source": "B",
                "Mapped": "old",
            },
            {
                "Name": "",
                "Text": "Narration",
                "Background": "",
                "Position": "center",
                "Voice": "",
                "Volume": "80",
                "Index": "old",
                "Source": "",
                "Mapped": "remove",
            },
        ]
    ).to_excel(path, sheet_name="Scenario", index=False)
    return path


def test_conditions_support_composition_and_missing_columns():
    row = pd.Series({"Name": "Alice", "Text": "Hello", "Voice": ""})

    assert matches_condition(
        row,
        {
            "all": [
                {"nonempty_any": ["Speaker", "Name"]},
                {"empty": "Voice"},
            ]
        },
    )
    assert not matches_condition(row, {"nonempty": "Missing"})
    assert matches_condition(row, {"not_in": {"Name": ["Bob"]}})


def test_definition_loads_all_builtin_operations(tmp_path):
    operation_names = [
        "set_default",
        "copy_nearest",
        "sequence",
        "template",
        "map_value",
        "clear_when",
        "copy_column",
    ]
    operations = []
    for index, operation_name in enumerate(operation_names):
        operation = {"id": f"op-{index}", "op": operation_name, "target": "Target"}
        if operation_name == "template":
            operation["value"] = "{counter}"
        if operation_name in ("map_value", "copy_column"):
            operation["source"] = "Source"
        if operation_name == "map_value":
            operation["mapping"] = {"A": "B"}
        operations.append(operation)

    definition = load_definition(write_config(tmp_path / "rules.yaml", operations))

    assert [operation.operation_type for operation in definition.operations] == operation_names


def test_definition_loads_project_plugin(tmp_path, monkeypatch):
    (tmp_path / "project_operation.py").write_text(
        "from core.table_transform import TableOperation\n"
        "class ProjectOperation(TableOperation):\n"
        "    operation_type = 'project_operation'\n"
        "    def plan(self, context):\n"
        "        return []\n",
        encoding="utf-8",
    )
    monkeypatch.syspath_prepend(str(tmp_path))
    config_path = write_config(
        tmp_path / "rules.yaml",
        [{"id": "project", "op": "project_operation", "target": "Target"}],
    )
    config = yaml.safe_load(config_path.read_text(encoding="utf-8"))
    config["plugins"] = ["project_operation:ProjectOperation"]
    config_path.write_text(
        yaml.safe_dump(config, allow_unicode=True, sort_keys=False),
        encoding="utf-8",
    )

    definition = load_definition(config_path)

    assert definition.operations[0].operation_type == "project_operation"


def test_dry_run_plans_without_modifying_workbook(tmp_path):
    workbook = write_workbook(tmp_path / "scenario.xlsx")
    config = write_config(
        tmp_path / "rules.yaml",
        [
            {
                "id": "voice",
                "op": "template",
                "target": "Voice",
                "value": "{speaker}_{counter:03d}.wav",
                "when": {"nonempty": ["Name", "Text"]},
                "options": {
                    "counter_scope": "file",
                    "counter_by": ["speaker"],
                    "variables": {"speaker": "Name"},
                },
            },
            {"id": "volume", "op": "set_default", "target": "Volume", "value": "100"},
            {
                "id": "background",
                "op": "copy_nearest",
                "target": "Background",
                "options": {"direction": "upward"},
            },
        ],
    )

    result = TableTransformEngine.from_config(config).process_file(workbook, dry_run=True)

    assert result.success
    assert not result.applied
    changes = {(change.row, change.column): change.new_value for change in result.plan.changes}
    assert changes[(2, "Voice")] == "Alice_001.wav"
    assert changes[(3, "Voice")] == "Alice_002.wav"
    assert changes[(3, "Background")] == "room"
    assert changes[(2, "Volume")] == "100"
    assert load_workbook(workbook)["Scenario"]["E2"].value is None


def test_apply_executes_sequence_mapping_copy_and_clear(tmp_path):
    workbook = write_workbook(tmp_path / "scenario.xlsx")
    config = write_config(
        tmp_path / "rules.yaml",
        [
            {
                "id": "index",
                "op": "sequence",
                "target": "Index",
                "when": {"nonempty": "Text"},
                "options": {"start": 10, "step": 5},
            },
            {
                "id": "mapping",
                "op": "map_value",
                "source": "Source",
                "target": "Mapped",
                "mapping": {"A": "Alpha", "B": "Beta"},
            },
            {
                "id": "copy",
                "op": "copy_column",
                "source": "Source",
                "target": "Voice",
            },
            {
                "id": "clear",
                "op": "clear_when",
                "target": "Mapped",
                "when": {"empty": "Name"},
            },
        ],
    )

    result = TableTransformEngine.from_config(config).process_file(workbook, dry_run=False)

    assert result.success
    assert result.applied
    sheet = load_workbook(workbook, data_only=False)["Scenario"]
    assert [sheet[f"G{row}"].value for row in (2, 3, 4)] == [10, 15, 20]
    assert [sheet[f"I{row}"].value for row in (2, 3, 4)] == ["Alpha", "Beta", None]
    assert [sheet[f"E{row}"].value for row in (2, 3, 4)] == ["A", "B", None]


def test_conflicting_operations_fail_without_writing(tmp_path):
    workbook = write_workbook(tmp_path / "scenario.xlsx")
    config = write_config(
        tmp_path / "rules.yaml",
        [
            {"id": "first", "op": "set_default", "target": "Volume", "value": "100"},
            {
                "id": "second",
                "op": "set_default",
                "target": "Volume",
                "value": "90",
                "options": {"overwrite": True},
            },
        ],
    )

    result = TableTransformEngine.from_config(config).process_file(workbook, dry_run=False)

    assert not result.success
    assert "同时被" in result.error
    assert load_workbook(workbook)["Scenario"]["F2"].value is None


def test_last_conflict_policy_keeps_original_value(tmp_path):
    workbook = write_workbook(tmp_path / "scenario.xlsx")
    config = write_config(
        tmp_path / "rules.yaml",
        [
            {"id": "first", "op": "set_default", "target": "Volume", "value": "100"},
            {
                "id": "second",
                "op": "set_default",
                "target": "Volume",
                "value": "90",
                "options": {"overwrite": True},
            },
        ],
        conflict_policy="last",
    )

    result = TableTransformEngine.from_config(config).process_file(workbook, dry_run=True)
    first_row = next(
        change for change in result.plan.changes if change.row == 2 and change.column == "Volume"
    )

    assert first_row.operation == "second"
    assert first_row.original_value == ""
    assert first_row.new_value == "90"
    assert result.plan.conflicts
